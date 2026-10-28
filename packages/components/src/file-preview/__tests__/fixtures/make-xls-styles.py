"""
生成老版 xls 样式检查用的夹具：边框（线型、颜色、空白单元格、合并区域）、填充、字体、对齐。

依赖：openpyxl、LibreOffice（soffice）。先用 openpyxl 写 xlsx，再由 LibreOffice 另存为 Excel 97-2003（BIFF8）。
用法：python make-xls-styles.py [输出目录]
- xls-styles.xls
"""
import os
import shutil
import subprocess
import sys
import tempfile

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

out_dir = sys.argv[1] if len(sys.argv) > 1 else os.path.dirname(os.path.abspath(__file__))

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Styles'
thin = Side(style='thin', color='000000')
grid = Border(left=thin, right=thin, top=thin, bottom=thin)
# A1:C3 细线网格（含空白单元格 C3）
for r in range(1, 4):
    for c in range(1, 4):
        cell = ws.cell(r, c)
        if (r, c) != (3, 3):
            cell.value = r * 10 + c
        cell.border = grid
# 线型与颜色
ws['E1'] = 'medium red'
ws['E1'].border = Border(bottom=Side(style='medium', color='FF0000'))
ws['E2'] = 'dashed'
ws['E2'].border = Border(top=Side(style='dashed', color='000000'))
ws['E3'] = 'double'
ws['E3'].border = Border(bottom=Side(style='double', color='000000'))
# 填充
ws['A5'] = 'solid'
ws['A5'].fill = PatternFill('solid', fgColor='FFFF00')
ws['B5'] = 'gray50'
ws['B5'].fill = PatternFill('mediumGray', fgColor='000000', bgColor='FFFFFF')
# 字体
ws['A6'] = 'bold'
ws['A6'].font = Font(name='Arial', bold=True)
ws['B6'] = 'italic red'
ws['B6'].font = Font(italic=True, color='FF0000')
ws['C6'] = 'underline'
ws['C6'].font = Font(underline='single', strike=True)
# 对齐
ws['A7'] = 'center top'
ws['A7'].alignment = Alignment(horizontal='center', vertical='top')
ws['B7'] = 'wrap text wrap text'
ws['B7'].alignment = Alignment(wrap_text=True)
# 合并区域：外框写在每个边缘单元格上（Excel 的保存方式）
ws.merge_cells('A9:C10')
ws['A9'] = 'merged'
for r in (9, 10):
    for c in (1, 2, 3):
        ws.cell(r, c).border = Border(
            left=thin if c == 1 else None,
            right=thin if c == 3 else None,
            top=thin if r == 9 else None,
            bottom=thin if r == 10 else None,
        )

tmp = tempfile.mkdtemp()
try:
    src = os.path.join(tmp, 'xls-styles.xlsx')
    wb.save(src)
    subprocess.run(['soffice', '--headless', '--convert-to', 'xls:MS Excel 97', '--outdir', tmp, src], check=True,
                   stdout=subprocess.DEVNULL)
    shutil.copy(os.path.join(tmp, 'xls-styles.xls'), os.path.join(out_dir, 'xls-styles.xls'))
finally:
    shutil.rmtree(tmp, ignore_errors=True)
