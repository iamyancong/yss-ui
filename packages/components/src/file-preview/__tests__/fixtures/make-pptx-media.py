"""
生成 pptx-media.pptx：每页一种图片情况，用于检查「每张图片都能显示，或显示带说明的占位，不出现裂图」。
内容全部是程序画的色块与英文标签，不含任何业务内容。

依赖：python-pptx、Pillow。用法：python make-pptx-media.py [输出路径]
"""
import io
import re
import struct
import sys
import zipfile

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Emu, Pt

OUT = sys.argv[1] if len(sys.argv) > 1 else 'pptx-media.pptx'
COLORS = [(22, 119, 255), (82, 196, 26), (250, 173, 20), (245, 34, 45), (114, 46, 209), (19, 194, 194)]


def card(index, size=(320, 180), mode='RGB'):
    """简单的色块图：边框 + 柱子，每张颜色 / 柱高不同（避免 python-pptx 按内容去重）"""
    color = COLORS[index % len(COLORS)]
    im = Image.new('RGB', size, (250, 250, 250))
    d = ImageDraw.Draw(im)
    d.rectangle([4, 4, size[0] - 5, size[1] - 5], outline=color, width=4)
    for i in range(5):
        h = 20 + ((i * 37 + index * 11) % 100)
        d.rectangle([24 + i * 56, size[1] - 20 - h, 60 + i * 56, size[1] - 20], fill=color)
    return im.convert(mode) if mode != 'RGB' else im


def encode(im, fmt, **kw):
    buf = io.BytesIO()
    im.save(buf, fmt, **kw)
    return buf.getvalue()


def emf_vector():
    """最小矢量 EMF：画刷 + 画笔 + 矩形 + 椭圆 + 三角形（不含位图，库无法显示）"""
    recs = []
    def rec(t, payload):
        recs.append(struct.pack('<II', t, 8 + len(payload)) + payload)
    rec(39, struct.pack('<IIII', 1, 0, 0x00FF7716, 0))          # CREATEBRUSHINDIRECT #1677ff
    rec(37, struct.pack('<I', 1))                                 # SELECTOBJECT
    rec(38, struct.pack('<IIiiI', 2, 0, 3, 0, 0x001F1F1F))       # CREATEPEN
    rec(37, struct.pack('<I', 2))
    rec(43, struct.pack('<iiii', 10, 10, 190, 110))              # RECTANGLE
    rec(39, struct.pack('<IIII', 3, 0, 0x001AC452, 0))
    rec(37, struct.pack('<I', 3))
    rec(42, struct.pack('<iiii', 210, 20, 310, 120))             # ELLIPSE
    pts = [(40, 170), (100, 125), (160, 170)]
    rec(86, struct.pack('<iiiiI', 40, 125, 160, 170, 3) + b''.join(struct.pack('<hh', *p) for p in pts))  # POLYGON16
    rec(14, struct.pack('<III', 0, 16, 20))                      # EOF
    body = b''.join(recs)
    return emf_header(320, 180, len(body), len(recs), 4) + body


def emf_header(w, h, body_len, body_records, handles):
    size = 108
    frame = (0, 0, int(w * 2645 / 100) , int(h * 2645 / 100))   # 0.01mm（96dpi）
    return struct.pack('<II4i4iIIIIHHIIIiiiiIIIii', 1, size, 0, 0, w - 1, h - 1, *frame, 0x464D4520, 0x10000,
                       size + body_len, body_records + 1, handles, 0, 0, 0, 0, 1920, 1080, 508, 286, 0, 0, 0, 508000, 286000)


def emf_bitmap(im):
    """只含一张 24 位位图（EMR_STRETCHDIBITS）的 EMF：库能取出位图显示"""
    w, h = im.size
    stride = (w * 3 + 3) & ~3
    rows = []
    px = im.convert('RGB').load()
    for y in range(h - 1, -1, -1):
        row = b''.join(bytes((px[x, y][2], px[x, y][1], px[x, y][0])) for x in range(w))
        rows.append(row + b'\0' * (stride - w * 3))
    bits = b''.join(rows)
    bmi = struct.pack('<IiiHHIIiiII', 40, w, h, 1, 24, 0, len(bits), 3780, 3780, 0, 0)
    head = 80
    payload = struct.pack('<4iiiiiiiIIIIIIii', 0, 0, w - 1, h - 1, 0, 0, 0, 0, w, h, head, len(bmi), head + len(bmi), len(bits), 0, 0x00CC0020, w, h)
    rec = struct.pack('<II', 81, 8 + len(payload) + len(bmi) + len(bits)) + payload + bmi + bits
    eof = struct.pack('<IIIII', 14, 20, 0, 16, 20)
    body = rec + eof
    return emf_header(w, h, len(body), 1, 1) + body


def wmf_vector():
    """最小可放置 WMF：矩形 + 椭圆"""
    recs = []
    def rec(fn, *words):
        recs.append(struct.pack('<IH', 3 + len(words), fn) + b''.join(struct.pack('<H', w & 0xFFFF) for w in words))
    rec(0x020B, 0, 0)                 # SETWINDOWORG y x
    rec(0x020C, 180, 320)             # SETWINDOWEXT y x
    rec(0x02FC, 0, 0x7716, 0x00FF, 0)  # CREATEBRUSHINDIRECT style, color lo, color hi, hatch
    rec(0x012D, 0)                    # SELECTOBJECT
    rec(0x041B, 110, 190, 10, 10)     # RECTANGLE bottom right top left
    rec(0x02FC, 0, 0xC452, 0x001A, 0)
    rec(0x012D, 1)
    rec(0x0418, 160, 300, 30, 200)    # ELLIPSE
    rec(0x0000)                       # EOF
    body = b''.join(recs)
    words = (18 + len(body)) // 2
    meta = struct.pack('<HHHIHIH', 1, 9, 0x0300, words, 2, max(len(r) // 2 for r in recs), 0)
    placeable = struct.pack('<IHhhhhHI', 0x9AC6CDD7, 0, 0, 0, 320, 180, 96, 0)
    checksum = 0
    for i in range(0, 20, 2):
        checksum ^= struct.unpack('<H', placeable[i:i + 2])[0]
    return placeable + struct.pack('<H', checksum) + meta + body


SVG = (b'<svg xmlns="http://www.w3.org/2000/svg" width="320" height="180" viewBox="0 0 320 180">'
       b'<rect x="4" y="4" width="312" height="172" rx="12" fill="#e6f4ff" stroke="#1677ff" stroke-width="4"/>'
       b'<circle cx="90" cy="90" r="50" fill="#52c41a"/><rect x="170" y="45" width="100" height="90" fill="#faad14"/></svg>')

HUGE_PNG = encode(card(30), 'PNG')
HUGE_PNG = HUGE_PNG[:16] + struct.pack('>II', 20000, 12000) + HUGE_PNG[24:]   # 头部声明 2.4 亿像素

# (标题, 替换后的媒体文件名, 数据, 对 XML 的额外处理)
CASES = [
    ('png', 'a.png', encode(card(0), 'PNG'), None),
    ('jpeg', 'b.jpeg', encode(card(1), 'JPEG', quality=85), None),
    ('gif', 'c.gif', encode(card(2).convert('P'), 'GIF'), None),
    ('bmp', 'd.bmp', encode(card(3), 'BMP'), None),
    ('tiff-lzw', 'e.tiff', encode(card(4), 'TIFF', compression='tiff_lzw'), None),
    ('tiff-raw', 'f.tif', encode(card(5), 'TIFF'), None),
    ('tiff-jpeg', 'g.tiff', encode(card(6), 'TIFF', compression='jpeg'), None),
    ('emf-vector', 'h.emf', emf_vector(), None),
    ('emf-bitmap', 'i.emf', emf_bitmap(card(7, (160, 90))), None),
    ('wmf', 'j.wmf', wmf_vector(), None),
    ('svg-with-png-fallback', 'k.png', encode(card(8), 'PNG'), 'svg'),
    ('cmyk-jpeg', 'l.jpg', encode(card(9, mode='CMYK'), 'JPEG', quality=85), None),
    ('external-link', 'm.png', encode(card(10), 'PNG'), 'external'),
    ('missing-part', 'n.png', None, 'missing'),
    ('missing-rel-id', 'o.png', encode(card(11), 'PNG'), 'norel'),
    ('cropped', 'p.png', encode(card(12), 'PNG'), 'crop'),
    ('huge', 'q.png', HUGE_PNG, None),
    ('jpeg-named-png', 'r.png', encode(card(13), 'JPEG'), None),
    ('svg-named-png', 's.png', SVG, None),
    ('unknown-format', 't.wdp', b'II\xbc\x01' + bytes(range(256)) * 4, None),
    ('empty-file', 'u.png', b'', None),
    ('corrupt-png', 'v.png', encode(card(14), 'PNG')[:60] + b'\0' * 200, None),
]


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    for i, (title, *_rest) in enumerate(CASES):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Emu(457200), Emu(228600), Emu(9000000), Emu(600000))
        box.text_frame.text = f'{i + 1}. {title}'
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
        src = io.BytesIO(encode(card(100 + i), 'PNG'))
        slide.shapes.add_picture(src, Emu(1828800), Emu(1143000), Emu(8534400), Emu(4800600))
    raw = io.BytesIO()
    prs.save(raw)
    zin = zipfile.ZipFile(io.BytesIO(raw.getvalue()))
    files = {n: zin.read(n) for n in zin.namelist()}

    for i, (title, name, data, extra) in enumerate(CASES):
        slide = f'ppt/slides/slide{i + 1}.xml'
        rels_path = f'ppt/slides/_rels/slide{i + 1}.xml.rels'
        rels = files[rels_path].decode()
        xml = files[slide].decode()
        m = re.search(r'Id="(rId\d+)"[^>]*Target="\.\./media/(image\d+\.png)"', rels) or re.search(
            r'Target="\.\./media/(image\d+\.png)"[^>]*Id="(rId\d+)"', rels)
        rid, old = (m.group(1), m.group(2)) if m.group(1).startswith('rId') else (m.group(2), m.group(1))
        files.pop(f'ppt/media/{old}', None)
        if data is not None:
            files[f'ppt/media/{name}'] = data
        rels = rels.replace(f'../media/{old}', f'../media/{name}')
        if extra == 'svg':
            files['ppt/media/k.svg'] = SVG
            rels = rels.replace('</Relationships>', '<Relationship Id="rIdSvg" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="../media/k.svg"/></Relationships>')
            xml = re.sub(r'<a:blip r:embed="(\w+)"\s*/>', r'<a:blip r:embed="\1"><a:extLst><a:ext uri="{96DAC541-7B7A-43D3-8B79-37D633B846F1}"><asvg:svgBlip xmlns:asvg="http://schemas.microsoft.com/office/drawing/2016/SVG/main" r:embed="rIdSvg"/></a:ext></a:extLst></a:blip>', xml)
        elif extra == 'external':
            files.pop(f'ppt/media/{name}', None)
            rels = rels.replace(f'Target="../media/{name}"', 'Target="https://example.com/remote-image.png" TargetMode="External"')
            xml = xml.replace(f'r:embed="{rid}"', f'r:link="{rid}"')
        elif extra == 'norel':
            xml = xml.replace(f'r:embed="{rid}"', 'r:embed="rId999"')
        elif extra == 'crop':
            xml = re.sub(r'(</a:blip>|<a:blip [^>]*/>)', r'\1<a:srcRect l="20000" t="10000" r="20000" b="10000"/>', xml, count=1)
        files[rels_path] = rels.encode()
        files[slide] = xml.encode()

    ct = files['[Content_Types].xml'].decode()
    for ext, mime in [('jpeg', 'image/jpeg'), ('jpg', 'image/jpeg'), ('gif', 'image/gif'), ('bmp', 'image/bmp'),
                      ('tiff', 'image/tiff'), ('tif', 'image/tiff'), ('emf', 'image/x-emf'), ('wmf', 'image/x-wmf'),
                      ('svg', 'image/svg+xml'), ('wdp', 'image/vnd.ms-photo')]:
        if f'Extension="{ext}"' not in ct:
            ct = ct.replace('<Default ', f'<Default Extension="{ext}" ContentType="{mime}"/><Default ', 1)
    files['[Content_Types].xml'] = ct.encode()

    with zipfile.ZipFile(OUT, 'w', zipfile.ZIP_DEFLATED) as z:
        for n in ['[Content_Types].xml'] + [n for n in files if n != '[Content_Types].xml']:
            info = zipfile.ZipInfo(n, date_time=(2026, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(info, files[n])
    print(OUT, len(CASES), 'slides')


build()
