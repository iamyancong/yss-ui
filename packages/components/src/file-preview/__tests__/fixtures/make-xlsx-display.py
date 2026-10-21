"""
生成表格显示检查用的 xlsx / xls：数字格式、公式结果、布局（合并、隐藏、冻结、换行、字号）、主题色、
没有保存计算结果的公式。内容全部是通用的示例文字与数字，不含任何业务内容。

依赖：xlsxwriter、openpyxl；可选 LibreOffice（soffice）生成 LibreOffice 另存版本、xls 版本与参考 CSV。
用法：python make-xlsx-display.py [输出目录]

生成的文件：
- xlsx-display.xlsx        xlsxwriter 写入，公式带缓存结果（去掉了 fullCalcOnLoad；xlsxwriter 无法写空字符串结果，
                           该情况由 LibreOffice 另存版本覆盖）
- xlsx-display-lo.xlsx     上面的文件经 LibreOffice 另存（LibreOffice 重新计算并保存结果）
- xlsx-display.xls         LibreOffice 另存为 Excel 97-2003（BIFF8）
- xlsx-display.expected.json  LibreOffice 按「显示的文字」导出的每个工作表内容（审计参考）
- xlsx-recalc.xlsx         xlsxwriter 默认写法：公式结果为 0 且要求打开时重算（fullCalcOnLoad）
- xlsx-no-cache.xlsx       openpyxl 写入：公式没有缓存结果；另含主题色与 tint
"""
import csv
import glob
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile

import openpyxl
import xlsxwriter
from openpyxl.styles import Font, PatternFill
from openpyxl.styles.colors import Color

OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.dirname(os.path.abspath(__file__))
D = 46287.0  # 2026-09-22

FORMATS = [
    ('General', 3), ('General', 0), ('General', 0.45), ('General', 1234567.891), ('General', 123456789012),
    ('General', -2.5), ('General', 1 / 3),
    ('0', 3.6), ('0.0', 3), ('0.00', -1.005), ('#,##0', 1234567.4), ('#,##0.00', -9876.543),
    ('0%', 0.45), ('0.0%', 0.4567), ('0.00E+00', 12345.678), ('##0.0E+0', 12345.678),
    ('"¥"#,##0.00', 1234.5), ('[$¥-804]#,##0.00', 1234.5), ('$#,##0.00_);($#,##0.00)', -1234.5),
    ('#,##0;[Red]-#,##0', -42), ('0.00;-0.00;"zero"', 0), ('0;"neg";"z";"t:"@', -3), ('@', 12.5),
    ('# ?/?', 0.75), ('?/8', 0.3), ('0.0,,"M"', 12345678), ('000-0000', 1234567),
    ('[>=100]"big";[<0]"neg";0', 150), ('[>=100]"big";[<0]"neg";0', 5),
    ('yyyy-mm-dd', D), ('yyyy\\-mm\\-dd', D), ('yyyy/m/d', D), ('m/d/yy', D), ('d-mmm-yy', D),
    ('yyyy"年"m"月"d"日"', D), ('m"月"d"日"', D), ('yyyy"年"m"月"', D),
    ('h:mm', D + 0.6), ('hh:mm:ss', 0.75), ('[h]:mm', 1.5), ('[mm]:ss', 0.0423), ('yyyy-mm-dd hh:mm', D + 0.25),
    ('h"时"mm"分"', 0.6), ('0.0', 'text'),
]


def build_display(path, remove_full_calc=True):
    wb = xlsxwriter.Workbook(path)
    bold = wb.add_format({'bold': True})
    ws = wb.add_worksheet('formats')
    ws.set_column(0, 0, 34)
    ws.set_column(1, 1, 22)
    ws.write_row(0, 0, ['format', 'value'], bold)
    for i, (code, value) in enumerate(FORMATS, start=1):
        fmt = wb.add_format({'num_format': code})
        ws.write_string(i, 0, code)
        if isinstance(value, str):
            ws.write_string(i, 1, value, fmt)
        else:
            ws.write_number(i, 1, value, fmt)

    fs = wb.add_worksheet('formulas')
    fs.set_column(0, 0, 26)
    fs.set_column(1, 1, 16)
    fs.write_row(0, 0, ['case', 'result'], bold)
    pct = wb.add_format({'num_format': '0%'})
    one = wb.add_format({'num_format': '0.0'})
    day = wb.add_format({'num_format': 'yyyy\\-mm\\-dd', 'align': 'center'})
    rows = [
        ('percent', '=0.9/2', pct, 0.45),
        ('one decimal', '=1+2', one, 3),
        ('zero', '=0*5', one, 0),
        ('date', '=DATE(2026,9,22)', day, D),
        ('text in date cell', '=IF(1>2,1,"—")', day, '—'),
        ('divide by zero', '=1/0', None, '#DIV/0!'),
        ('not available', '=NA()', None, '#N/A'),
        ('boolean', '=1>0', None, True),
        ('sum', '=SUM(B2:B4)', one, 3.45),
    ]
    for i, (label, formula, fmt, value) in enumerate(rows, start=1):
        fs.write_string(i, 0, label)
        fs.write_formula(i, 1, formula, fmt, value)
    fs.write_string(11, 0, 'plain bool')
    fs.write_boolean(11, 1, False)
    fs.write_string(12, 0, 'rich text')
    fs.write_rich_string(12, 1, 'plain ', bold, 'bold', ' tail')
    fs.write_string(13, 0, 'link')
    fs.write_url(13, 1, 'https://example.com/', string='example link')
    fs.activate()

    lay = wb.add_worksheet('layout')
    title = wb.add_format({'bold': True, 'font_size': 14, 'align': 'center', 'valign': 'vcenter'})
    head = wb.add_format({'bold': True, 'bg_color': '#2F5597', 'font_color': '#FFFFFF', 'text_wrap': True,
                          'align': 'center', 'valign': 'vcenter', 'border': 1})
    lay.merge_range('A1:F1', 'Merged title across six columns', title)
    lay.set_row(0, 30)
    lay.write_row(1, 0, ['Name', 'A long header that should wrap inside the cell', 'Hidden column', 'Amount', 'Ratio', 'Note'], head)
    lay.set_row(1, 36)
    lay.set_column(0, 0, 12)
    lay.set_column(1, 1, 14)
    lay.set_column(2, 2, 10, None, {'hidden': True})
    lay.set_column(3, 4, 11)
    lay.set_column(5, 5, 24)
    money = wb.add_format({'num_format': '#,##0.00'})
    ratio = wb.add_format({'num_format': '0.0%'})
    styles = [
        wb.add_format({'font_size': 9}), wb.add_format({'font_size': 11, 'italic': True}),
        wb.add_format({'font_size': 14, 'underline': 1}), wb.add_format({'font_size': 18, 'font_strikeout': True}),
    ]
    for i in range(8):
        r = i + 2
        lay.write_string(r, 0, f'Item {i + 1}')
        lay.write_string(r, 1, 'x' * (i + 1))
        lay.write_string(r, 2, 'secret')
        lay.write_number(r, 3, 1000 * (i + 1) + 0.5, money)
        lay.write_number(r, 4, (i + 1) / 10, ratio)
        lay.write_string(r, 5, f'size sample {i}', styles[i % 4])
    lay.set_row(5, None, None, {'hidden': True})
    lay.freeze_panes(2, 1)
    wrap = wb.add_format({'text_wrap': True})
    lay.write_string(10, 5, 'This cell has wrapped text without a stored row height so it needs more lines', wrap)

    hid = wb.add_worksheet('hidden sheet')
    hid.write_string(0, 0, 'should not be shown')
    hid.hide()
    wb.close()
    if remove_full_calc:
        strip_full_calc(path)


def strip_full_calc(path):
    tmp = path + '.tmp'
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == 'xl/workbook.xml':
                data = re.sub(rb'\s*fullCalcOnLoad="1"', b'', data)
            dst.writestr(item, data)
    os.replace(tmp, path)


def build_recalc(path):
    wb = xlsxwriter.Workbook(path)
    ws = wb.add_worksheet('recalc')
    ws.write_row(0, 0, ['a', 'b', 'sum'])
    ws.write_row(1, 0, [1, 2])
    ws.write_formula(1, 2, '=A2+B2')
    wb.close()


def build_no_cache(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'no cache'
    ws.append(['a', 'b', 'sum', 'ratio'])
    ws.append([1, 2, '=A2+B2', '=A2/B2'])
    ws['D2'].number_format = '0%'
    th = wb.create_sheet('theme')
    for i, tint in enumerate([0, 0.4, 0.8, -0.25, -0.5]):
        cell = th.cell(row=i + 1, column=1, value=f'accent1 tint {tint}')
        cell.fill = PatternFill('solid', fgColor=Color(theme=4, tint=tint))
        font_cell = th.cell(row=i + 1, column=2, value=f'text dk2 tint {tint}')
        font_cell.font = Font(color=Color(theme=3, tint=tint))
    wb.save(path)


def soffice(args):
    exe = shutil.which('soffice') or shutil.which('libreoffice')
    if not exe:
        return False
    subprocess.run([exe, '--headless', *args], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    return True


def export_expected(xlsx, out_json):
    """LibreOffice 按「显示的文字」导出所有工作表（隐藏表也导出，审计时忽略）"""
    tmp = tempfile.mkdtemp()
    src = os.path.join(tmp, 'book.xlsx')
    shutil.copy(xlsx, src)
    if not soffice(['--convert-to', 'csv:Text - txt - csv (StarCalc):44,34,76,1,,0,false,true,true,false,false,-1', '--outdir', tmp, src]):
        return
    sheets = {}
    for f in sorted(glob.glob(os.path.join(tmp, 'book-*.csv'))):
        name = os.path.basename(f)[len('book-'):-len('.csv')]
        with io.open(f, encoding='utf-8') as fh:
            sheets[name] = [row for row in csv.reader(fh)]
    with io.open(out_json, 'w', encoding='utf-8') as fh:
        json.dump(sheets, fh, ensure_ascii=False, indent=1)
    shutil.rmtree(tmp)


def main():
    os.makedirs(OUT, exist_ok=True)
    display = os.path.join(OUT, 'xlsx-display.xlsx')
    build_display(display)
    build_recalc(os.path.join(OUT, 'xlsx-recalc.xlsx'))
    build_no_cache(os.path.join(OUT, 'xlsx-no-cache.xlsx'))
    tmp = tempfile.mkdtemp()
    work = os.path.join(tmp, 'xlsx-display-lo.xlsx')
    shutil.copy(display, work)
    lo_dir = os.path.join(tmp, 'lo')
    if soffice(['--convert-to', 'xlsx:Calc MS Excel 2007 XML', '--outdir', lo_dir, work]):
        shutil.copy(os.path.join(lo_dir, 'xlsx-display-lo.xlsx'), os.path.join(OUT, 'xlsx-display-lo.xlsx'))
        soffice(['--convert-to', 'xls:MS Excel 97', '--outdir', lo_dir, display])
        shutil.copy(os.path.join(lo_dir, 'xlsx-display.xls'), os.path.join(OUT, 'xlsx-display.xls'))
        export_expected(display, os.path.join(OUT, 'xlsx-display.expected.json'))
    shutil.rmtree(tmp)


if __name__ == '__main__':
    main()
